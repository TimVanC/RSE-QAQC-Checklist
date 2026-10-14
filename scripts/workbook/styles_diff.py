import sys
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

path = sys.argv[1]
wb = load_workbook(path, data_only=False)
ws = wb["Task1"]

print("=== SHEET PROTECTION Task1 ===", "sheet protected:", ws.protection.sheet)
print("=== COLUMN WIDTHS ===")
for k, d in sorted(ws.column_dimensions.items()):
    print(f"  {k}: width={d.width} hidden={d.hidden}")
print("=== ROW DIMS (non-default) ===")
for r, d in sorted(ws.row_dimensions.items()):
    if d.height is not None or d.hidden:
        print(f"  row {r}: height={d.height} hidden={d.hidden}")
print("=== CONDITIONAL FORMATTING ===")
for cf in ws.conditional_formatting:
    for rule in cf.rules:
        print(f"  {cf.sqref}: type={rule.type} operator={rule.operator} formula={rule.formula} dxf_fill={getattr(rule.dxf,'fill',None) and rule.dxf.fill.fgColor.rgb if rule.dxf and rule.dxf.fill else None}")
print("=== DATA VALIDATIONS ===")
for d in ws.data_validations.dataValidation:
    print(f"  {d.sqref} type={d.type} formula1={d.formula1!r}")

def desc(c):
    f = c.fill
    fill = f.fgColor.rgb if f and f.fill_type else None
    b = c.border
    bs = "".join(x[0].upper() if getattr(b, x).style else "." for x in ("left","right","top","bottom"))
    return f"fill={fill} border={bs} locked={c.protection.locked} fmt={c.number_format!r} halign={c.alignment.horizontal} font_b={c.font.b} color={c.font.color.rgb if c.font.color else None}"

print("=== COLUMN L/M/N STYLES rows 9-76 ===")
for r in range(9, 77):
    a = ws.cell(r, 1).value
    parts = []
    for col in ("L", "M", "N"):
        c = ws[f"{col}{r}"]
        parts.append(f"{col}: {desc(c)}")
    print(f"  row {r:2} step={a!r:5} | " + " || ".join(parts))

print("=== HEADER / DATE / PARTS INPUT CELL STYLES ===")
for coord in ["C3","C4","C5","C6","C7","D3","E3","C26","D26","E26","C27","C28","M26","M27","M28","C34","D34","E34","F34","H34","I34","J34","M34","E78","E79","E80","C68","C69","C70","E70","J70","K70","L70","M70"]:
    c = ws[coord]
    print(f"  {coord:4} value={c.value!r:30} {desc(c)}")

print("=== ROW 9-64 full-row style scan: which cells (A..N) have a fill or an all-sides border ===")
for r in range(9, 65):
    hits = []
    for ci in range(1, 15):
        c = ws.cell(r, ci)
        f = c.fill
        fill = f.fgColor.rgb if f and f.fill_type else None
        b = c.border
        full = all(getattr(b, x).style for x in ("left","right","top","bottom"))
        if fill or full:
            hits.append(f"{get_column_letter(ci)}({'F' if fill else ''}{'B' if full else ''}{':'+str(fill) if fill else ''})")
    print(f"  row {r:2}: {' '.join(hits)}")

# DIFF all tabs vs their template
print("=== DIFF Task2..35 vs Task1, CALLBACK2..10 vs CALLBACK1 (values+formulas) ===")
def cells(wsx):
    d = {}
    for row in wsx.iter_rows():
        for c in row:
            if c.value is not None:
                d[c.coordinate] = c.value
    return d
base_t = cells(wb["Task1"]); base_c = cells(wb["CALLBACK1"])
for name in wb.sheetnames:
    if name.startswith("Task") and name != "Task1":
        base = base_t
    elif name.startswith("CALLBACK") and name != "CALLBACK1":
        base = base_c
    else:
        continue
    cur = cells(wb[name])
    diffs = [(k, base.get(k), cur.get(k)) for k in sorted(set(base)|set(cur), key=lambda s:(int(''.join(ch for ch in s if ch.isdigit())), s)) if base.get(k) != cur.get(k)]
    print(f"  {name}: {'IDENTICAL' if not diffs else diffs}")
print("=== DIFF Task1 vs CALLBACK1 ===")
cur = base_c
for k in sorted(set(base_t)|set(base_c), key=lambda s:(int(''.join(ch for ch in s if ch.isdigit())), s)):
    if base_t.get(k) != base_c.get(k):
        print(f"  {k}: Task1={base_t.get(k)!r}  CALLBACK1={base_c.get(k)!r}")
# style diff Task1 vs CALLBACK1 for column M
print("=== STYLE DIFF col M Task1 vs CALLBACK1 ===")
wc = wb["CALLBACK1"]
for r in range(9, 77):
    if desc(ws[f"M{r}"]) != desc(wc[f"M{r}"]):
        print(f"  M{r}: Task1={desc(ws[f'M{r}'])} | CALLBACK1={desc(wc[f'M{r}'])}")
print("(end)")
