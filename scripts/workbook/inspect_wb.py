import sys, json
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

path = sys.argv[1]
wb_f = load_workbook(path, data_only=False)   # formulas
wb_v = load_workbook(path, data_only=True)    # cached values

print("=== SHEETS ===")
for ws in wb_f.worksheets:
    print(f"{ws.title!r:40} dims={ws.dimensions:14} max_row={ws.max_row} max_col={ws.max_column} state={ws.sheet_state}")
print()

def dump_sheet(name, col_limit=None, row_limit=None):
    ws = wb_f[name]; wv = wb_v[name]
    print(f"=== SHEET {name!r} ===")
    print("merged:", sorted([str(r) for r in ws.merged_cells.ranges], key=lambda s:(int(''.join(c for c in s.split(':')[0] if c.isdigit())), s)))
    for row in ws.iter_rows(min_row=1, max_row=row_limit or ws.max_row, max_col=col_limit or ws.max_column):
        for c in row:
            if c.value is None: continue
            v = c.value
            cached = wv[c.coordinate].value
            if isinstance(v, str) and v.startswith("="):
                print(f"  {c.coordinate:6} FORMULA {v!r:70} cached={cached!r}")
            else:
                s = repr(v)
                if len(s) > 160: s = s[:160] + "..."
                print(f"  {c.coordinate:6} {s}")
    # data validations
    dv = ws.data_validations.dataValidation if ws.data_validations else []
    if dv:
        print("  data validations:")
        for d in dv:
            print(f"    {d.sqref} type={d.type} formula1={d.formula1!r} allow_blank={d.allow_blank}")
    print()

for name in ["Employee Production (Monthly)", "JobLog", "Task1", "CALLBACK1"]:
    if name in wb_f.sheetnames:
        dump_sheet(name)
    else:
        print("MISSING SHEET", name)
