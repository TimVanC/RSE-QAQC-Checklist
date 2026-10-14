import sys
from openpyxl import load_workbook
wb = load_workbook(sys.argv[1])

def fillinfo(c):
    f = c.fill
    if not f or not f.fill_type: return "-"
    fg = f.fgColor
    if fg.type == "theme": return f"theme{fg.theme}{'%+.2f'%fg.tint if fg.tint else ''}"
    if fg.type == "rgb": return f"rgb{fg.rgb}"
    return f"{fg.type}{fg.value}"

for sheet in ["Task1", "CALLBACK1", "Employee Production (Monthly)", "JobLog"]:
    ws = wb[sheet]
    print(f"=== {sheet}: protected={ws.protection.sheet} UNLOCKED (input) cells ===")
    for row in ws.iter_rows():
        for c in row:
            if not c.protection.locked:
                step = ws.cell(c.row, 1).value
                label = ws.cell(c.row, 2).value if c.column > 2 else None
                print(f"  {c.coordinate:5} fmt={c.number_format:10} fill={fillinfo(c):14} step={step!r:5} rowtext={str(label)[:60]!r}")
    print()

ws = wb["Task1"]
print("=== Task1 fill colors on key rows (B and M) ===")
for r in [10, 11, 21, 22, 24, 25, 26, 29, 30, 34, 35, 36, 43, 47, 53, 54, 60, 61, 64]:
    print(f"  row {r:2}: A={fillinfo(ws.cell(r,1)):14} B={fillinfo(ws.cell(r,2)):14} M={fillinfo(ws.cell(r,13)):14}")
print("=== Task1 conditional formatting dxf details ===")
for cf in ws.conditional_formatting:
    for rule in cf.rules:
        d = rule.dxf
        fill = None; font = None
        if d is not None:
            if d.fill is not None:
                fg = d.fill.fgColor; bg = d.fill.bgColor
                fill = f"fg={fg.rgb if fg is not None and fg.type=='rgb' else (fg.type if fg is not None else None)} bg={bg.rgb if bg is not None and bg.type=='rgb' else (bg.type if bg is not None else None)}"
            if d.font is not None and d.font.color is not None:
                font = d.font.color.rgb if d.font.color.type == 'rgb' else d.font.color.type
        cs = None
        if rule.colorScale is not None:
            cs = [(v.type, v.val) for v in rule.colorScale.cfvo], [c.rgb for c in rule.colorScale.color]
        print(f"  {str(cf.sqref):14} {rule.type:14} {rule.operator or '':14} {rule.formula} fill[{fill}] font[{font}] cs={cs}")
print("=== Task1 column N / row 1-8 contents ===")
for row in ws.iter_rows(min_col=14, max_col=14):
    for c in row:
        if c.value is not None: print("  N:", c.coordinate, repr(c.value))
print("  print area:", ws.print_area, " freeze:", ws.freeze_panes)
print("=== Workbook defined names ===")
try:
    for n, dn in wb.defined_names.items(): print("  ", n, dn.attr_text)
except Exception as e:
    print("  (none/err)", e)
print("=== Employee Production: labels with adjacent cells ===")
wp = wb["Employee Production (Monthly)"]
for r in range(1, 27):
    vals = [(c.coordinate, c.value) for c in wp[r] if c.value is not None]
    if vals: print("  ", vals)
